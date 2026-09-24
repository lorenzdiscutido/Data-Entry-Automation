import os
import time
import json
import datetime
import threading
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from watchdog.observers.polling import PollingObserver 
from watchdog.events import FileSystemEventHandler

from config import INCOMING_FOLDER, PROCESSED_LOG, REFERENCE_DATA
from image_utils import load_image, prepare_excel_image
from gemini_client import get_raw_response, parse_and_clean_json

last_processed_time = time.time()
is_processing = False

def load_processed_files():
    if not os.path.exists(PROCESSED_LOG):
        return set()
    with open(PROCESSED_LOG, 'r', encoding='utf-8') as f:
        return set(line.strip() for line in f)

def mark_as_processed(filepath):
    with open(PROCESSED_LOG, 'a', encoding='utf-8') as f:
        f.write(f"{filepath}\n")

def _populate_reference_sheet(wb):
    ws_ref = wb.create_sheet(title="Reference Data")
    header_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    
    for row_idx, row_values in enumerate(REFERENCE_DATA, start=1):
        ws_ref.append(row_values)
        for col_idx in range(1, 6):
            cell = ws_ref.cell(row=row_idx, column=col_idx)
            if row_idx == 1:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    ws_ref.column_dimensions['A'].width = 15
    ws_ref.column_dimensions['B'].width = 26
    ws_ref.column_dimensions['C'].width = 24
    ws_ref.column_dimensions['D'].width = 45
    ws_ref.column_dimensions['E'].width = 40
    ws_ref.row_dimensions[1].height = 28

def process_single_image(filepath):
    global last_processed_time, is_processing
    processed_files = load_processed_files()
    if filepath in processed_files:
        return False 
        
    try:
        is_processing = True
        print(f"Processing natively via Google Gemini API: {os.path.basename(filepath)}")
        
        img = load_image(filepath)
        raw_text = get_raw_response(img)
        parsed_data = parse_and_clean_json(raw_text)
        
        rel_path = os.path.relpath(os.path.dirname(filepath), start=INCOMING_FOLDER)
        if rel_path != ".":
            folder_parts = rel_path.replace('\\', '/').split('/')
            for part in folder_parts:
                part_lower = part.lower()
                
                if "tower" in part_lower and not parsed_data.get("Tower"):
                    parsed_data["Tower"] = part
                elif "drop" in part_lower and not parsed_data.get("Drop"):
                    parsed_data["Drop"] = part
                elif any(x in part_lower for x in ["elevation", "front", "back", "left", "right"]) and not parsed_data.get("Elevation"):
                    parsed_data["Elevation"] = part

        current_month_year = datetime.datetime.now().strftime("%B_%Y") 
        output_xlsx_path = f"master_output_{current_month_year}.xlsx"

        max_retries = 5
        for attempt in range(max_retries):
            try:
                if os.path.exists(output_xlsx_path):
                    wb = openpyxl.load_workbook(output_xlsx_path)
                    ws = wb["Whiteboard Data"] if "Whiteboard Data" in wb.sheetnames else wb.active
                    if "Reference Data" not in wb.sheetnames:
                        _populate_reference_sheet(wb)
                else:
                    wb = openpyxl.Workbook()
                    ws = wb.active
                    ws.title = "Whiteboard Data"
                    
                    ws.append([
                        "Submitter", "Date", "Elevation", "Drop", "Floor", "Tower",
                        "Defects", "", "Whiteboard Photo",
                        "POSSIBLE CAUSE", "Possible Cause Justification", "Recommended Repair", "FINDINGS"
                    ])
                    ws.append([
                        "", "", "", "", "", "",
                        "Damage", "Dimension", "",
                        "", "", "", ""
                    ])

                    ws.merge_cells('G1:H1') 
                    
                    for col in ['A', 'B', 'C', 'D', 'E', 'F', 'I', 'J', 'K', 'L', 'M']:
                        ws.merge_cells(f'{col}1:{col}2')
                    
                    gray_fill = PatternFill(start_color="BFBFBF", end_color="BFBFBF", fill_type="solid")
                    
                    for row in ws['A1':'M2']:
                        for cell in row:
                            cell.font = Font(bold=True)
                            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

                    ws['M1'].fill = gray_fill
                    ws['M2'].fill = gray_fill

                    ws.column_dimensions['J'].width = 24
                    ws.column_dimensions['K'].width = 30
                    ws.column_dimensions['L'].width = 28
                    ws.column_dimensions['M'].width = 22

                    _populate_reference_sheet(wb)

                start_row = ws.max_row + 1
                
                # Memorize all valid defect codes from the Reference Data (e.g., 'BP', 'CC+', 'DS')
                VALID_CODES = {row[0] for row in REFERENCE_DATA[1:]}

                rows_data = []
                for mat in ["Sealant", "Concrete", "Paint", "Gasket"]:
                    rows_data.append((mat, "", True, False)) 
                    
                    dmg_list = parsed_data.get(f"{mat} Damage", [])
                    dim_list = parsed_data.get(f"{mat} Dimension", [])
                    
                    expanded_dmg = []
                    expanded_dim = []
                    
                    # Smart Split Logic
                    for i in range(max(len(dmg_list), len(dim_list))):
                        dmg_str = dmg_list[i] if i < len(dmg_list) else ""
                        dim_str = dim_list[i] if i < len(dim_list) else ""
                        
                        tokens = str(dmg_str).split()
                        found_codes = [t for t in tokens if t in VALID_CODES]
                        
                        # If we find multiple valid codes in one cell (e.g., "BP FP")
                        if len(found_codes) > 1:
                            for code in found_codes:
                                expanded_dmg.append(code)
                                expanded_dim.append(dim_str) # Duplicates the dimension for each defect
                        else:
                            # Keep it intact (e.g., "DS C-C" or regular single codes)
                            expanded_dmg.append(dmg_str)
                            expanded_dim.append(dim_str)
                    
                    max_len = max(len(expanded_dmg), 1) 
                    
                    for i in range(max_len):
                        dmg_val = expanded_dmg[i] if i < len(expanded_dmg) else ""
                        dim_val = expanded_dim[i] if i < len(expanded_dim) else ""
                        rows_data.append((dmg_val, dim_val, False, True)) 

                total_rows = len(rows_data)

                static_keys = ["Submitter", "Date", "Elevation", "Drop", "Floor", "Tower"]
                for i, key in enumerate(static_keys):
                    ws.cell(row=start_row, column=i+1).value = parsed_data.get(key, "")
                    if total_rows > 1:
                        ws.merge_cells(start_row=start_row, start_column=i+1, end_row=start_row + total_rows - 1, end_column=i+1)

                for i, (g_val, h_val, is_title, is_data) in enumerate(rows_data):
                    r = start_row + i
                    cell_g = ws.cell(row=r, column=7)
                    cell_h = ws.cell(row=r, column=8)
                    
                    cell_g.value = g_val
                    cell_h.value = h_val

                    if is_title:
                        ws.merge_cells(start_row=r, start_column=7, end_row=r, end_column=8)
                        cell_g.font = Font(bold=True)
                    elif is_data:
                        lookup_val = f'LEFT($G{r}, FIND(" ", $G{r}&" ") - 1)'
                        
                        ws.cell(row=r, column=10).value = f'=IFERROR(VLOOKUP({lookup_val}, \'Reference Data\'!$A$2:$E$14, 3, FALSE), "")'
                        ws.cell(row=r, column=11).value = f'=IFERROR(VLOOKUP({lookup_val}, \'Reference Data\'!$A$2:$E$14, 4, FALSE), "")'
                        ws.cell(row=r, column=12).value = f'=IFERROR(VLOOKUP({lookup_val}, \'Reference Data\'!$A$2:$E$14, 5, FALSE), "")'
                        ws.cell(row=r, column=13).value = f'=IFERROR(VLOOKUP({lookup_val}, \'Reference Data\'!$A$2:$E$14, 2, FALSE), "")'

                photo_col_letter = 'I' 
                excel_img = prepare_excel_image(filepath, total_rows)
                ws.add_image(excel_img, f"{photo_col_letter}{start_row}")
                
                if total_rows > 1:
                    ws.merge_cells(f"{photo_col_letter}{start_row}:{photo_col_letter}{start_row + total_rows - 1}")
                
                base_height = max(110 / total_rows, 20)
                for i in range(total_rows):
                    r = start_row + i
                    ws.row_dimensions[r].height = base_height
                    for c in range(1, 14):
                        ws.cell(row=r, column=c).alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                ws.column_dimensions[photo_col_letter].width = 25

                wb.save(output_xlsx_path)
                break 
                
            except PermissionError:
                print(f"[Warning] {output_xlsx_path} is open. Retrying in 10s... ({attempt + 1}/{max_retries})")
                time.sleep(10)
        else:
            print(f"[Error] Could not write to Excel after {max_retries} attempts.")
            return False
            
        mark_as_processed(filepath)
        print(f"[Success] Appended fully automated record to {output_xlsx_path}.\n")
        
        last_processed_time = time.time()
        time.sleep(4) 
        return True
        
    except json.JSONDecodeError:
        print(f"[Failed] The AI did not return a valid JSON format. Raw output: {raw_text}\n")
    except Exception as e:
        print(f"[Failed] Error processing file {filepath}: {e}\n")
    finally:
        last_processed_time = time.time()
    return False

def check_batch_completion():
    global last_processed_time, is_processing
    while True:
        if is_processing and (time.time() - last_processed_time > 8):
            print("[Status] All new uploaded files finish. Continuing to monitor for new files...")
            is_processing = False
        time.sleep(2)

def run_batch_processor():
    print(f"\n--- Starting Batch Process ---")
    print(f"Scanning: {INCOMING_FOLDER}")
    processed_count = 0
    for root, dirs, files in os.walk(INCOMING_FOLDER):
        for filename in files:
            if filename.lower().endswith(('.png', '.jpg', '.jpeg')):
                filepath = os.path.join(root, filename)
                if process_single_image(filepath):
                    processed_count += 1
    print(f"--- Batch Process Complete. Processed {processed_count} new files. ---\n")

class PhotoHandler(FileSystemEventHandler):
    def trigger_processing(self, filepath):
        if os.path.basename(filepath).startswith('.'):
            return
        if not filepath.lower().endswith(('.png', '.jpg', '.jpeg')):
            return
            
        print(f"\n[Detected] New file synced: {filepath}")
        time.sleep(3) 
        process_single_image(filepath)

    def on_created(self, event):
        if event.is_directory:
            return
        self.trigger_processing(event.src_path)

    def on_moved(self, event):
        if event.is_directory:
            return
        self.trigger_processing(event.dest_path)

def run_watchdog():
    completion_thread = threading.Thread(target=check_batch_completion, daemon=True)
    completion_thread.start()

    event_handler = PhotoHandler()
    observer = PollingObserver() 
    observer.schedule(event_handler, path=INCOMING_FOLDER, recursive=True)
    observer.start()
    print(f"[Status] Watchdog Started. Actively polling for new files in: {INCOMING_FOLDER}")
    print("Press Ctrl+C to stop the watcher anytime.\n")
    try:
        while True:
            time.sleep(1)
    except KeyboardInterrupt:
        observer.stop()
        print("\n[Status] Stopped folder watcher.")
    observer.join()

if __name__ == "__main__":
    if not os.path.exists(INCOMING_FOLDER):
        print(f"Error: The folder {INCOMING_FOLDER} does not exist.")
    else:
        print("========================================")
        print(" WHITEBOARD DATA EXTRACTION SCRIPT")
        print(" (Dynamic Split Update)")
        print("========================================")
        print("1: Run Batch Process Only")
        print("2: Start Watchdog Only (Live Monitor)")
        print("3: Catch Up (Batch) THEN Start Watchdog")
        print("========================================")
        
        choice = input("Enter your choice (1, 2, or 3): ").strip()
        
        if choice == "1":
            run_batch_processor()
        elif choice == "2":
            run_watchdog()
        elif choice == "3":
            run_batch_processor()
            run_watchdog()
        else:
            print("Invalid choice. Exiting script.") 